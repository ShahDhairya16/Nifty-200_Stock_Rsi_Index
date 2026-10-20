from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, Optional

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from app.database.repositories import StockRepository, PriceRepository, RSIRepository
from app.services.stock_universe_service import StockUniverseService
from app.utils.date_utils import format_date_iso, parse_date
from app.utils.logger import logger


BASE_DIR = Path(__file__).resolve().parent.parent.parent
OUTPUT_DIR = BASE_DIR / "output"


class ExcelReportService:
    """Generate Excel reports from MongoDB market data."""

    @staticmethod
    def generate_nifty200_excel_report(
        output_dir: Optional[Path] = None,
        refresh_market_data: bool = True,
        end_date: Optional[date] = None,
        in_memory: bool = False,
    ) -> Dict[str, Any]:

        report_end = parse_date(end_date) if end_date is not None else date.today()

        # ------------------------------------------------------------------
        # 1. Synchronize the current NIFTY 200 universe.
        #
        # The stock-universe collection is the single source of truth for
        # which constituents are currently active.
        #
        # We intentionally synchronize the universe before reading active
        # stocks so that newly added/deactivated constituents are reflected
        # in the report immediately.
        # ------------------------------------------------------------------
        universe_summary = StockUniverseService.sync_nifty200_universe()

        if universe_summary.get("errors"):
            error_message = (
                "NIFTY 200 universe synchronization failed: "
                + "; ".join(str(error) for error in universe_summary["errors"])
            )
            logger.error(error_message)
            raise RuntimeError(error_message)

        logger.info(
            "NIFTY 200 universe synchronized successfully. "
            f"Fetched={universe_summary.get('total_fetched', 0)}, "
            f"New={universe_summary.get('new_stocks', 0)}, "
            f"Updated={universe_summary.get('updated_stocks', 0)}, "
            f"Deactivated={universe_summary.get('deactivated_stocks', 0)}"
        )

        # ------------------------------------------------------------------
        # 2. Load the current active NIFTY 200 universe.
        #
        # This is the authoritative stock list for the complete report.
        # No stock list is derived from historical price/RSI collections.
        # ------------------------------------------------------------------
        stocks = StockRepository.get_active_stocks()

        if not stocks:
            raise ValueError(
                "No active NIFTY 200 stocks found in MongoDB after "
                "universe synchronization."
            )

        active_symbols = {
            str(stock["symbol"]).strip().upper()
            for stock in stocks
            if stock.get("symbol")
        }

        if not active_symbols:
            raise ValueError(
                "Active NIFTY 200 universe is empty after symbol normalization."
            )

        symbols = sorted(active_symbols)

        logger.info(
            f"Active NIFTY 200 count: {len(active_symbols)}"
        )

        # ------------------------------------------------------------------
        # 3. Load historical prices.
        #
        # Historical price data may contain deactivated NIFTY constituents.
        # Therefore, explicitly filter the price dataframe against the
        # current active universe before creating the price matrix.
        # ------------------------------------------------------------------
        prices = PriceRepository.load_prices(end_date=report_end)

        if prices.empty:
            raise ValueError(
                "No market data found in MongoDB. "
                "Run --sync-universe and --backfill-historical first."
            )

        prices["symbol"] = (
            prices["symbol"]
            .astype(str)
            .str.strip()
            .str.upper()
        )

        # IMPORTANT:
        # Historical/deactivated symbols must never leak into the report.
        prices = prices[
            prices["symbol"].isin(active_symbols)
        ].copy()

        if prices.empty:
            raise ValueError(
                "No price data found for the current active NIFTY 200 universe."
            )

        # ------------------------------------------------------------------
        # 4. Select the last 70 available trading days.
        #
        # This preserves the existing report behavior.
        # ------------------------------------------------------------------
        dates = sorted(prices["trade_date"].unique())[-70:]

        if not dates:
            raise ValueError(
                "No trading dates found for the active NIFTY 200 universe."
            )

        prices = prices[
            prices["trade_date"].isin(dates)
        ].copy()

        # ------------------------------------------------------------------
        # 5. Create the closing-price matrix using ONLY active symbols.
        #
        # reindex(columns=symbols) ensures that an active constituent with
        # missing price data remains represented as a blank/NaN column
        # rather than being replaced by an old constituent.
        # ------------------------------------------------------------------
        matrix = (
            prices
            .pivot_table(
                index="trade_date",
                columns="symbol",
                values="close_price",
            )
            .reindex(columns=symbols)
            .reset_index()
        )

        # ------------------------------------------------------------------
        # 6. Load RSI data.
        #
        # The RSI collection can contain historical records for stocks that
        # are no longer active. Explicitly filter them out.
        # ------------------------------------------------------------------
        rsi = RSIRepository.load_rsi(end_date=report_end)

        if rsi.empty:
            rsi = pd.DataFrame(
                columns=[
                    "symbol",
                    "trade_date",
                    "rsi_22",
                    "rsi_44",
                    "rsi_66",
                    "average_rsi",
                ]
            )
        else:
            rsi["symbol"] = (
                rsi["symbol"]
                .astype(str)
                .str.strip()
                .str.upper()
            )

            rsi = rsi[
                rsi["symbol"].isin(active_symbols)
            ].copy()

        # ------------------------------------------------------------------
        # 7. Get the latest RSI record for each active symbol.
        # ------------------------------------------------------------------
        if not rsi.empty:
            latest = (
                rsi
                .sort_values(["symbol", "trade_date"])
                .groupby("symbol", as_index=False)
                .tail(1)
                .copy()
            )
        else:
            latest = pd.DataFrame(
                columns=[
                    "symbol",
                    "trade_date",
                    "rsi_22",
                    "rsi_44",
                    "rsi_66",
                    "average_rsi",
                ]
            )

        # ------------------------------------------------------------------
        # 8. Build RSI ranking from the ACTIVE UNIVERSE first.
        #
        # This is important.
        #
        # We do NOT build the ranking directly from the RSI collection
        # because a newly added active stock may temporarily have no RSI
        # record.
        #
        # Instead:
        #
        #       Active NIFTY 200
        #              ↓
        #         LEFT JOIN RSI
        #
        # This guarantees that an active constituent remains in the report
        # even if its RSI/price data is temporarily unavailable.
        # ------------------------------------------------------------------
        universe_df = pd.DataFrame(
            [
                {
                    "symbol": str(stock["symbol"]).strip().upper(),
                    "company_name": (
                        stock.get("company_name")
                        or str(stock["symbol"]).strip().upper()
                    ),
                }
                for stock in stocks
                if stock.get("symbol")
            ]
        )

        ranking = universe_df.merge(
            latest,
            on="symbol",
            how="left",
        )

        # ------------------------------------------------------------------
        # 9. Sort RSI ranking.
        #
        # Existing RSI ranking behavior is preserved:
        # highest average RSI first.
        #
        # Stocks without RSI data are placed at the bottom instead of being
        # removed from the active universe.
        # ------------------------------------------------------------------
        ranking = (
            ranking
            .sort_values(
                "average_rsi",
                ascending=False,
                na_position="last",
            )
            .reset_index(drop=True)
        )

        ranking.insert(
            0,
            "Rank",
            range(1, len(ranking) + 1),
        )

        ranking = ranking.rename(
            columns={
                "trade_date": "Latest Date",
            }
        )[
            [
                "Rank",
                "symbol",
                "company_name",
                "Latest Date",
                "rsi_22",
                "rsi_44",
                "rsi_66",
                "average_rsi",
            ]
        ]

        ranking.columns = [
            "Rank",
            "Symbol",
            "Company Name",
            "Latest Date",
            "RSI 22",
            "RSI 44",
            "RSI 66",
            "Average RSI",
        ]

        # ------------------------------------------------------------------
        # 10. Validate the universe BEFORE creating/saving the workbook.
        #
        # Both stock-level Excel sheets must contain exactly the same active
        # NIFTY 200 universe.
        # ------------------------------------------------------------------

        # Closing Price sheet symbols
        closing_price_symbols = {
            str(symbol).strip().upper()
            for symbol in symbols
        }

        # RSI Ranking sheet symbols
        rsi_ranking_symbols = {
            str(symbol).strip().upper()
            for symbol in ranking["Symbol"].dropna()
        }

        # Validate closing-price sheet
        missing_from_closing_price = (
            active_symbols - closing_price_symbols
        )

        unexpected_in_closing_price = (
            closing_price_symbols - active_symbols
        )

        # Validate RSI Ranking sheet
        missing_from_rsi = (
            active_symbols - rsi_ranking_symbols
        )

        unexpected_from_rsi = (
            rsi_ranking_symbols - active_symbols
        )

        logger.info(
            "NIFTY 200 Excel universe validation:"
        )

        logger.info(
            f"Active NIFTY 200 count: {len(active_symbols)}"
        )

        logger.info(
            f"Closing Price Excel stock count: "
            f"{len(closing_price_symbols)}"
        )

        logger.info(
            f"RSI Ranking Excel stock count: "
            f"{len(rsi_ranking_symbols)}"
        )

        logger.info(
            f"Missing from Closing Price: "
            f"{sorted(missing_from_closing_price)}"
        )

        logger.info(
            f"Unexpected in Closing Price: "
            f"{sorted(unexpected_in_closing_price)}"
        )

        logger.info(
            f"Missing from RSI Ranking: "
            f"{sorted(missing_from_rsi)}"
        )

        logger.info(
            f"Unexpected in RSI Ranking: "
            f"{sorted(unexpected_from_rsi)}"
        )

        # ------------------------------------------------------------------
        # 11. Fail loudly if either stock-level sheet does not exactly match
        #     the active NIFTY 200 universe.
        #
        # This prevents a wrong Excel report from being silently generated.
        # ------------------------------------------------------------------
        if (
            missing_from_closing_price
            or unexpected_in_closing_price
            or missing_from_rsi
            or unexpected_from_rsi
        ):
            raise ValueError(
                "NIFTY 200 Excel universe validation failed. "
                f"Active={len(active_symbols)}, "
                f"ClosingPrice={len(closing_price_symbols)}, "
                f"RSIRanking={len(rsi_ranking_symbols)}, "
                f"MissingClosingPrice={sorted(missing_from_closing_price)}, "
                f"UnexpectedClosingPrice={sorted(unexpected_in_closing_price)}, "
                f"MissingRSI={sorted(missing_from_rsi)}, "
                f"UnexpectedRSI={sorted(unexpected_from_rsi)}"
            )

        logger.info(
            "NIFTY 200 Excel universe validation PASSED. "
            "Active universe exactly matches both stock-level sheets."
        )

        # ------------------------------------------------------------------
        # 12. Create workbook.
        #
        # Existing workbook layout, sheet names, columns and formatting are
        # preserved.
        # ------------------------------------------------------------------
        workbook = openpyxl.Workbook()

        sheet = workbook.active
        sheet.title = "Last 70 Days Closing Price"

        sheet.append(["Date"] + symbols)

        for row in matrix.itertuples(index=False, name=None):
            trade_date = row[0]

            if hasattr(trade_date, "strftime"):
                formatted_date = trade_date.strftime("%d-%b-%Y")
            else:
                formatted_date = str(trade_date)

            sheet.append(
                [formatted_date] + list(row[1:])
            )

        rank_sheet = workbook.create_sheet(
            "RSI Ranking"
        )

        rank_sheet.append(
            list(ranking.columns)
        )

        for row in ranking.itertuples(
            index=False,
            name=None,
        ):
            rank_sheet.append(
                list(row)
            )

        # ------------------------------------------------------------------
        # 13. Existing Excel formatting.
        # ------------------------------------------------------------------
        header_font = Font(
            name="Segoe UI",
            bold=True,
            color="FFFFFF",
        )

        header_fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )

        for ws in workbook.worksheets:
            ws.freeze_panes = (
                "B2"
                if ws == sheet
                else "A2"
            )

            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center"
                )

            ws.auto_filter.ref = ws.dimensions

            for column in ws.columns:
                ws.column_dimensions[
                    column[0].column_letter
                ].width = min(
                    max(
                        max(
                            len(str(cell.value or ""))
                            for cell in column
                        ) + 2,
                        12,
                    ),
                    38,
                )

        # ------------------------------------------------------------------
        # 14. Save workbook.
        # ------------------------------------------------------------------
        filename = (
            f"NIFTY200_RSI_Report_"
            f"{report_end.isoformat()}_"
            f"{datetime.now():%H%M%S}.xlsx"
        )

        if in_memory:
            buffer = BytesIO()
            workbook.save(buffer)

            report_data = buffer.getvalue()
            path = None

        else:
            target_dir = output_dir or OUTPUT_DIR
            target_dir.mkdir(
                parents=True,
                exist_ok=True,
            )

            path = target_dir / filename

            workbook.save(path)

            report_data = None

        # ------------------------------------------------------------------
        # 15. Return existing response structure plus validation information.
        # ------------------------------------------------------------------
        return {
            "success": True,
            "file_path": str(path) if path else None,
            "data": report_data,
            "filename": filename,
            "latest_trade_date": format_date_iso(max(dates)),
            "trading_days": len(dates),
            "stocks": len(symbols),
            "active_nifty200_count": len(active_symbols),
            "closing_price_stock_count": len(closing_price_symbols),
            "rsi_ranking_stock_count": len(rsi_ranking_symbols),
            "missing_stocks": sorted(
                active_symbols
                - closing_price_symbols
            ),
            "unexpected_stocks": sorted(
                closing_price_symbols
                - active_symbols
            ),
            "rsi_missing_stocks": sorted(
                active_symbols
                - rsi_ranking_symbols
            ),
            "rsi_unexpected_stocks": sorted(
                rsi_ranking_symbols
                - active_symbols
            ),
        }